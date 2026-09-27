"""Create synthetic original exports; no user files are read."""
from pathlib import Path
import sys,json,subprocess,os
import openpyxl,pandas as pd
from import_pipeline import export_zip,HERE
root=Path(sys.argv[1]);exports=root/'combined/payhere/raw/exports';exports.mkdir(parents=True,exist_ok=True)
for folder in ['combined/analysis','combined/metadata','payhere/normalized']:(root/folder).mkdir(parents=True,exist_ok=True)
months=[]
for start,end in [('2025-10-01','2025-12-31'),('2026-01-01','2026-04-30')]:
    wb=openpyxl.Workbook();wb.remove(wb.active);s=wb.create_sheet('매출 내역')
    rows=[[None]*18 for _ in range(13)];rows[3][0]=start+'~'+end;rows[8][1]=3000;rows[8][2]=1;rows[8][4]=0;rows[8][5]=0;rows[10][0]='결제일';rows[10][14]='환불 일시';rows[12]=[start,'10:00:00','예시 쿠키',3000,0,0,3000,0,0,0,0,0,0,0,0,0,0,0]
    for row in rows:s.append(row)
    s=wb.create_sheet('상품별');rows=[[None]*10 for _ in range(11)];rows[8][4]=1;rows[8][6]=3000;rows[8][8]=0;rows[8][9]=3000;rows[10]=[1,'디저트','예시 쿠키',3000,1,0,3000,0,0,3000]
    for row in rows:s.append(row)
    wb.save(exports/(start.replace('-','')+'~'+end.replace('-','')+'_synthetic.xlsx'))
    months.append(dict(month=start[:7],menu_original='예시 쿠키',ui_order_count=1,payment_amount_krw=3000))
pd.DataFrame(months).to_csv(root/'payhere/normalized/menu_monthly_ui.csv',index=False)
pd.DataFrame([dict(date=d,displayed_sales_krw=3000 if d in ['2025-10-01','2026-01-01'] else 0) for d in pd.date_range('2025-10-01','2026-04-30').strftime('%Y-%m-%d')]).to_csv(root/'payhere/normalized/daily_calendar_ui.csv',index=False)
archive=root/'baseline.zip';export_zip(archive,'2026-05-01','2026-05-03',[('2026-05-01',2),('2026-05-02',3),('2026-05-03',1)])
def run(script,args):
    r=subprocess.run([sys.executable,str(HERE/'import'/script),*args],env={**os.environ,'TOSS_EXPORT_PASSWORD':'synthetic'},capture_output=True)
    if r.returncode:raise RuntimeError(r.stderr.decode())
run('process_sales.py',['--zip',str(archive),'--out',str(root),'--captured-at','2026-05-03T12:00:00+09:00'])
run('build_data.py',['--root',str(root),'--out',str(root/'combined')])
export_zip(root/'next.zip','2026-05-02','2026-05-04',[('2026-05-02',4),('2026-05-03',1),('2026-05-04',5)])
