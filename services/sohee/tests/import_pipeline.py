"""Synthetic Toss exports: replacement, corrections, rollback guards and forecasts."""
import io,json,os,subprocess,sys,tempfile,unittest,zipfile,importlib.util
from pathlib import Path
import pandas as pd
import openpyxl
HERE=Path(__file__).resolve().parents[1]
spec=importlib.util.spec_from_file_location('combined',HERE/'import/build_data.py');combined=importlib.util.module_from_spec(spec);spec.loader.exec_module(combined)
def export_zip(path,start,end,values,attribution='결제/환불한 날'):
    wb=openpyxl.Workbook();wb.remove(wb.active)
    base=wb.create_sheet('데이터 기준');base.append(['기준']);base.append(['시작','끝','집계','시작시각','단위']);base.append([start,end,attribution,'00:00','일별'])
    payment=wb.create_sheet('결제 상세내역');payment.append(['테스트']);payment.append(['date','paid','channel','order','count','amount','vat','method','acquirer','status','cancelled'])
    item=wb.create_sheet('상품 주문 상세내역');item.append(['테스트']);item.append(list(range(19)))
    pday=wb.create_sheet('결제 합계');pday.append(['테스트']);pday.append(list(range(12)))
    iday=wb.create_sheet('상품 주문 합계');iday.append(list(range(10)))
    for index,(day,qty) in enumerate(values):
        amount=qty*3000;stamp=day+' 10:30:00';status='완료' if qty>=0 else '취소'
        payment.append([day,stamp,'매장',str(index),1 if qty>=0 else -1,amount,0,'카드','예시카드','승인' if qty>=0 else '취소',None])
        item.append([day,status,stamp,'매장',str(index),'예시 쿠키','test','디저트','','','',qty,amount,0,0,0,amount,amount,0])
        pday.append([day,amount,0,1 if qty>=0 else -1,None,0,amount,0,0,0,0,None])
        iday.append([day,'예시 쿠키','test','디저트',qty,amount,0,0,amount,0])
    stream=io.BytesIO();wb.save(stream)
    with zipfile.ZipFile(path,'w',zipfile.ZIP_DEFLATED) as z:z.writestr('synthetic.xlsx',stream.getvalue())
class Imports(unittest.TestCase):
    def test_replace_overlap_preserves_refunds_and_outside_range(self):
        with tempfile.TemporaryDirectory() as temp:
            root=Path(temp)/'private'; first=Path(temp)/'first.zip'; second=Path(temp)/'second.zip'
            export_zip(first,'2026-05-01','2026-05-03',[('2026-05-01',2),('2026-05-02',3),('2026-05-03',1)])
            export_zip(second,'2026-05-02','2026-05-04',[('2026-05-02',4),('2026-05-03',-1),('2026-05-04',5)])
            def run(z,captured):return subprocess.run([sys.executable,str(HERE/'import/process_sales.py'),'--zip',str(z),'--out',str(root),'--captured-at',captured],env={**os.environ,'TOSS_EXPORT_PASSWORD':'synthetic'},capture_output=True)
            one=run(first,'2026-05-03T12:00:00+09:00');self.assertEqual(one.returncode,0,one.stderr.decode())
            two=run(second,'2026-05-04T12:00:00+09:00');self.assertEqual(two.returncode,0,two.stderr.decode())
            data=pd.read_csv(root/'normalized/item_lines.csv');self.assertEqual(len(data),4);self.assertEqual(data.quantity.tolist(),[2,4,-1,5]);self.assertEqual(data.net_amount.sum(),30000)
            self.assertEqual(data.loc[data.partial_day,'date'].tolist(),['2026-05-04'])
            before=(root/'normalized/item_lines.csv').read_bytes();old=run(first,'2026-05-03T12:00:00+09:00');self.assertNotEqual(old.returncode,0);self.assertEqual(before,(root/'normalized/item_lines.csv').read_bytes())
            gap=Path(temp)/'gap.zip';export_zip(gap,'2026-05-10','2026-05-10',[('2026-05-10',2)]);self.assertNotEqual(run(gap,'2026-05-11T12:00:00+09:00').returncode,0);self.assertEqual(before,(root/'normalized/item_lines.csv').read_bytes())
            changed=Path(temp)/'changed.zip';export_zip(changed,'2026-05-02','2026-05-04',[('2026-05-02',4)],'주문한 날');self.assertNotEqual(run(changed,'2026-05-05T12:00:00+09:00').returncode,0)
    def test_partial_bulk_and_future_do_not_enter_preparation(self):
        dates=pd.date_range('2026-03-01','2026-04-30').strftime('%Y-%m-%d')
        items=pd.DataFrame([dict(date=d,partial_day=False,order_key=d,category='디저트',menu_original='예시 4구',quantity=2,net_amount=20000,hour=10,weekday=pd.Timestamp(d).weekday()) for d in dates])
        pay=pd.DataFrame([dict(date=d,partial_day=False,amount=20000) for d in dates])
        a=combined.forecast(items,pay,'2026-04-15')
        # Mutating future data cannot change any prior prediction or backtest.
        altered=items.copy();altered.loc[altered.date>'2026-04-15','quantity']=999
        b=combined.forecast(altered,pay,'2026-04-15');self.assertEqual(a[0],b[0]);pd.testing.assert_frame_equal(a[1],b[1])
        items.loc[items.date=='2026-04-01',['partial_day','quantity']]=[True,999];pay.loc[pay.date=='2026-04-01','partial_day']=True
        items.loc[items.date=='2026-04-02',['net_amount','quantity']]=[200000,50]
        plans,back,metrics,observed,bulk=combined.forecast(items,pay,'2026-04-15')
        self.assertNotIn('2026-04-01',observed.date.tolist());self.assertNotIn('2026-04-02',observed.date.tolist());self.assertTrue(all(p['unit']=='팩(4개)' for p in plans));self.assertTrue((back.train_end<back.date).all());self.assertTrue(all(p['end']<='2026-04-15' for p in plans))
if __name__=='__main__':unittest.main()
